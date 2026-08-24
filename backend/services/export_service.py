"""
NEW FILE. Excel export - spec section 12 ("Reports & Export System" -
PDF export already existed via services/report_generator.py, this adds
Excel export alongside it).

Deliberately generated ON DEMAND rather than pre-built and stored on disk
like the PDF report: no new DB table/columns, no new files cluttering the
reports/ folder, and it can never go stale relative to the analysis data
(always reflects exactly what's in Postgres right now). Reuses the exact
same dict shape as GET /analysis/{id} (services.analysis_service.
build_analysis_response), so there's a single source of truth for what
"the analysis data" means across the JSON API, the PDF report, and this.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)


def _autosize(ws, min_width=12, max_width=60):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, min_width), max_width)


def _write_header_row(ws, row_idx, headers):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def build_excel_report(analysis_data: dict) -> io.BytesIO:
    """
    analysis_data: the exact dict returned by
    services.analysis_service.build_analysis_response(). Builds a 4-sheet
    workbook and returns it as an in-memory buffer ready to stream back
    in a FastAPI response - nothing touches disk.
    """
    wb = Workbook()

    # ---- Sheet 1: Summary ----
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "SportsAI - Injury Risk Analysis Report"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    risk_summary = analysis_data.get("risk_score_summary") or {}
    summary_rows = [
        ("Analysis ID", analysis_data.get("analysis_id")),
        ("Athlete ID", analysis_data.get("athlete_id")),
        ("Video Filename", analysis_data.get("filename")),
        ("Status", analysis_data.get("status")),
        ("Overall Risk Score", f"{risk_summary.get('overall_score', '—')}%"),
        ("Overall Risk Level", risk_summary.get("risk_level", "—")),
    ]
    row = 3
    for label, value in summary_rows:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1
    _autosize(ws)

    # ---- Sheet 2: Injury Risks ----
    ws2 = wb.create_sheet("Injury Risks")
    _write_header_row(ws2, 1, ["Category", "Risk Level", "Probability (%)", "Reasons"])
    injury_risks = analysis_data.get("injury_risks") or {}
    row = 2
    for category, details in injury_risks.items():
        details = details or {}
        reasons = details.get("reasons") or []
        ws2.cell(row=row, column=1, value=category)
        ws2.cell(row=row, column=2, value=details.get("risk_level"))
        ws2.cell(row=row, column=3, value=details.get("probability"))
        ws2.cell(row=row, column=4, value="; ".join(reasons) if isinstance(reasons, list) else str(reasons))
        row += 1
    _autosize(ws2)

    # ---- Sheet 3: Biomechanics ----
    ws3 = wb.create_sheet("Biomechanics")
    _write_header_row(ws3, 1, ["Metric", "Value"])
    biomechanics = analysis_data.get("biomechanics") or {}
    row = 2

    def _flatten(prefix, value):
        nonlocal row
        if isinstance(value, dict):
            for k, v in value.items():
                _flatten(f"{prefix} - {k}" if prefix else k, v)
        else:
            ws3.cell(row=row, column=1, value=prefix)
            ws3.cell(row=row, column=2, value=str(value))
            row += 1

    _flatten("", biomechanics)
    _autosize(ws3)

    # ---- Sheet 4: Recommendations ----
    ws4 = wb.create_sheet("Recommendations")
    _write_header_row(ws4, 1, ["#", "Recommendation"])
    recommendations = analysis_data.get("recommendations") or []
    for idx, rec in enumerate(recommendations, start=1):
        ws4.cell(row=idx + 1, column=1, value=idx)
        ws4.cell(row=idx + 1, column=2, value=str(rec))
    _autosize(ws4)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer